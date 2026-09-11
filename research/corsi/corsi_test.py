from flask import Flask, render_template_string, jsonify, request, send_file
import json
import csv
import os
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

# Crear carpetas para datos
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, 'data')
os.makedirs(os.path.join(DATA_DIR, 'sessions'), exist_ok=True)
os.makedirs(os.path.join(DATA_DIR, 'movements'), exist_ok=True)
os.makedirs(os.path.join(DATA_DIR, 'exports'), exist_ok=True)

HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Corsi ML Data Collector - v3.3</title>
    <style>
        * { margin: 0; padding: 0; box-sizing: border-box; }
        
        body {
            font-family: 'Segoe UI', system-ui, sans-serif;
            background: linear-gradient(135deg, #1a1c2e 0%, #4a192c 100%);
            min-height: 100vh;
            color: #fff;
            padding: 20px;
        }

        .container { max-width: 1200px; margin: 0 auto; }

        .header {
            text-align: center;
            padding: 30px;
            background: rgba(255,255,255,0.05);
            border-radius: 20px;
            margin-bottom: 30px;
            backdrop-filter: blur(10px);
        }

        h1 { 
            font-size: 2.5rem; 
            background: linear-gradient(to right, #00d4ff, #7b2cbf);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
            margin-bottom: 10px;
        }

        .mode-selector {
            display: flex;
            gap: 20px;
            justify-content: center;
            margin: 20px 0;
            flex-wrap: wrap;
        }

        .mode-card {
            background: rgba(255,255,255,0.05);
            border: 2px solid rgba(255,255,255,0.1);
            border-radius: 15px;
            padding: 25px;
            cursor: pointer;
            transition: all 0.3s;
            width: 280px;
            text-align: center;
        }

        .mode-card:hover {
            border-color: #00d4ff;
            transform: translateY(-5px);
            background: rgba(0,212,255,0.1);
        }

        .mode-card.selected {
            border-color: #00d4ff;
            background: rgba(0,212,255,0.2);
            box-shadow: 0 0 30px rgba(0,212,255,0.3);
        }

        .mode-icon {
            font-size: 3rem;
            margin-bottom: 15px;
        }

        .mode-title {
            font-size: 1.3rem;
            font-weight: bold;
            margin-bottom: 10px;
            color: #00d4ff;
        }

        .mode-desc {
            font-size: 0.9rem;
            opacity: 0.8;
            line-height: 1.4;
        }

        .setup-panel, .game-panel, .results-panel {
            background: rgba(255,255,255,0.05);
            border-radius: 20px;
            padding: 30px;
            margin-bottom: 20px;
            backdrop-filter: blur(10px);
            border: 1px solid rgba(255,255,255,0.1);
        }

        .form-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(250px, 1fr));
            gap: 20px;
            margin-bottom: 20px;
        }

        .form-group {
            display: flex;
            flex-direction: column;
        }

        label {
            margin-bottom: 8px;
            color: #00d4ff;
            font-weight: 600;
        }

        input, select, textarea {
            padding: 12px;
            border-radius: 10px;
            border: 2px solid rgba(255,255,255,0.1);
            background: rgba(0,0,0,0.3);
            color: #fff;
            font-size: 1rem;
        }

        input:focus, select:focus, textarea:focus {
            outline: none;
            border-color: #00d4ff;
        }

        .metrics-display {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(150px, 1fr));
            gap: 15px;
            margin-bottom: 20px;
        }

        .metric-card {
            background: rgba(0,0,0,0.3);
            padding: 15px;
            border-radius: 15px;
            text-align: center;
            border-left: 4px solid #00d4ff;
        }

        .metric-label {
            font-size: 0.85rem;
            opacity: 0.7;
            margin-bottom: 5px;
        }

        .metric-value {
            font-size: 1.5rem;
            font-weight: bold;
            color: #00d4ff;
        }

        .corsi-board {
            position: relative;
            width: 100%;
            height: 600px;
            background: linear-gradient(145deg, #0f1419 0%, #1a2332 100%);
            border-radius: 20px;
            margin: 30px 0;
            box-shadow: inset 0 4px 20px rgba(0,0,0,0.8);
            overflow: hidden;
            border: 2px solid rgba(0, 212, 255, 0.2);
        }

        .cube {
            position: absolute;
            width: 80px;
            height: 80px;
            background: linear-gradient(145deg, #2d3748 0%, #1a202c 100%);
            border-radius: 16px;
            cursor: pointer;
            transition: all 0.2s cubic-bezier(0.4, 0, 0.2, 1);
            box-shadow: 
                0 10px 25px rgba(0,0,0,0.5),
                inset 0 2px 4px rgba(255,255,255,0.1);
            display: flex;
            align-items: center;
            justify-content: center;
            font-weight: bold;
            font-size: 1.5rem;
            color: rgba(255,255,255,0.4);
            border: 2px solid rgba(0, 212, 255, 0.1);
            user-select: none;
        }

        .cube:hover:not(.disabled):not(.active):not(.selected) {
            transform: translateY(-3px) scale(1.05);
            border-color: rgba(0, 212, 255, 0.5);
            box-shadow: 0 15px 35px rgba(0, 212, 255, 0.2);
        }

        .cube.active {
            background: linear-gradient(145deg, #00d4ff 0%, #0099cc 100%);
            box-shadow: 
                0 0 40px rgba(0, 212, 255, 0.8),
                0 0 80px rgba(0, 212, 255, 0.4),
                inset 0 2px 4px rgba(255,255,255,0.5);
            color: #000;
            transform: scale(1.15);
            border-color: #fff;
            z-index: 100;
        }

        .cube.selected {
            background: linear-gradient(145deg, #10b981 0%, #059669 100%);
            box-shadow: 0 0 30px rgba(16, 185, 129, 0.6);
            color: white;
            border-color: #34d399;
        }

        .cube.error {
            background: linear-gradient(145deg, #ef4444 0%, #dc2626 100%);
            animation: shake 0.5s;
            box-shadow: 0 0 30px rgba(239, 68, 68, 0.8);
        }

        .cube.disabled {
            opacity: 0.5;
            cursor: not-allowed;
        }

        @keyframes shake {
            0%, 100% { transform: translateX(0); }
            25% { transform: translateX(-10px); }
            75% { transform: translateX(10px); }
        }

        .controls {
            display: flex;
            gap: 15px;
            justify-content: center;
            flex-wrap: wrap;
        }

        button {
            padding: 15px 30px;
            border: none;
            border-radius: 12px;
            font-size: 1rem;
            font-weight: 600;
            cursor: pointer;
            transition: all 0.3s;
            text-transform: uppercase;
            letter-spacing: 1px;
        }

        .btn-primary {
            background: linear-gradient(145deg, #00d4ff 0%, #0099cc 100%);
            color: #000;
            box-shadow: 0 4px 15px rgba(0, 212, 255, 0.4);
        }

        .btn-success {
            background: linear-gradient(145deg, #10b981 0%, #059669 100%);
            color: white;
        }

        .btn-warning {
            background: linear-gradient(145deg, #f59e0b 0%, #d97706 100%);
            color: white;
        }

        .btn-secondary {
            background: rgba(255,255,255,0.1);
            color: white;
            border: 2px solid rgba(255,255,255,0.2);
        }

        button:hover {
            transform: translateY(-2px);
            box-shadow: 0 8px 25px rgba(0,0,0,0.3);
        }

        .status-bar {
            display: flex;
            justify-content: space-between;
            align-items: center;
            padding: 20px;
            background: rgba(0,0,0,0.3);
            border-radius: 15px;
            margin-bottom: 20px;
            flex-wrap: wrap;
            gap: 15px;
        }

        .mode-indicator {
            display: inline-block;
            padding: 8px 16px;
            border-radius: 20px;
            font-size: 0.9rem;
            font-weight: bold;
            margin-bottom: 15px;
        }

        .mode-normal {
            background: rgba(0, 212, 255, 0.2);
            border: 2px solid #00d4ff;
            color: #00d4ff;
        }

        .mode-inverse {
            background: rgba(245, 158, 11, 0.2);
            border: 2px solid #f59e0b;
            color: #f59e0b;
        }

        .message {
            padding: 20px;
            border-radius: 15px;
            text-align: center;
            font-size: 1.1rem;
            margin: 20px 0;
            min-height: 60px;
            display: flex;
            align-items: center;
            justify-content: center;
        }

        .message.info { background: rgba(0, 212, 255, 0.1); border: 2px solid rgba(0, 212, 255, 0.3); }
        .message.success { background: rgba(16, 185, 129, 0.1); border: 2px solid rgba(16, 185, 129, 0.3); }
        .message.error { background: rgba(239, 68, 68, 0.1); border: 2px solid rgba(239, 68, 68, 0.3); }

        .hidden { display: none !important; }

        .download-section {
            background: rgba(0, 212, 255, 0.05);
            border: 2px dashed rgba(0, 212, 255, 0.3);
            border-radius: 15px;
            padding: 30px;
            text-align: center;
            margin-top: 20px;
        }

        .feature-list {
            text-align: left;
            margin: 20px 0;
            padding-left: 20px;
        }

        .feature-list li {
            margin-bottom: 10px;
            color: rgba(255,255,255,0.8);
        }

        .feature-list li:before {
            content: "✓ ";
            color: #00d4ff;
            font-weight: bold;
        }

        .progress-bar {
            width: 100%;
            height: 6px;
            background: rgba(255,255,255,0.1);
            border-radius: 3px;
            overflow: hidden;
            margin-top: 10px;
        }

        .progress-fill {
            height: 100%;
            background: linear-gradient(90deg, #00d4ff, #7b2cbf);
            transition: width 0.3s;
        }

        .file-list {
            background: rgba(0,0,0,0.2);
            border-radius: 10px;
            padding: 15px;
            margin-top: 15px;
            max-height: 200px;
            overflow-y: auto;
        }

        .file-item {
            display: flex;
            justify-content: space-between;
            align-items: center;
            padding: 8px;
            border-bottom: 1px solid rgba(255,255,255,0.1);
        }

        @media (max-width: 768px) {
            .corsi-board { height: 400px; }
            .cube { width: 60px; height: 60px; font-size: 1.2rem; }
            h1 { font-size: 1.8rem; }
            .mode-card { width: 100%; }
        }
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <h1>🧠 Corsi ML Data Collector</h1>
            <p>Sistema avanzado de recolección para Machine Learning</p>
            <span style="background: #10b981; color: #000; padding: 5px 15px; border-radius: 20px; font-size: 0.9rem; font-weight: bold;">📁 v3.3 - Bug Fix Final</span>
        </div>

        <div class="setup-panel" id="modePanel">
            <h2 style="margin-bottom: 20px; color: #00d4ff; text-align: center;">Selecciona el Modo de Prueba</h2>
            <div class="mode-selector">
                <div class="mode-card" id="modeNormal" onclick="selectMode('normal')">
                    <div class="mode-icon">➡️</div>
                    <div class="mode-title">Corsi Directo</div>
                    <div class="mode-desc">Repite la secuencia en el <strong>mismo orden</strong></div>
                </div>
                <div class="mode-card" id="modeInverse" onclick="selectMode('inverse')">
                    <div class="mode-icon">🔄</div>
                    <div class="mode-title">Corsi Inverso</div>
                    <div class="mode-desc">Repite la secuencia en <strong>orden inverso</strong></div>
                </div>
            </div>
        </div>

        <div class="setup-panel hidden" id="setupPanel">
            <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 20px;">
                <h2 style="color: #00d4ff;">Configuración</h2>
                <button class="btn-secondary" onclick="backToMode()" style="padding: 8px 16px;">← Cambiar Modo</button>
            </div>
            <div id="modeDisplay" class="mode-indicator" style="display: none;"></div>
            <div class="form-grid">
                <div class="form-group">
                    <label>ID Participante *</label>
                    <input type="text" id="userId" placeholder="Ej: PAC-001" required>
                </div>
                <div class="form-group">
                    <label>Edad *</label>
                    <input type="number" id="age" min="5" max="100" required>
                </div>
                <div class="form-group">
                    <label>Género</label>
                    <select id="gender">
                        <option value="M">Masculino</option>
                        <option value="F">Femenino</option>
                        <option value="O">Otro</option>
                    </select>
                </div>
                <div class="form-group">
                    <label>Condición</label>
                    <select id="condition">
                        <option value="control">Control</option>
                        <option value="adhd">TDAH</option>
                        <option value="mci">Deterioro Cognitivo</option>
                        <option value="alzheimer">Alzheimer</option>
                        <option value="tdc">Traumatismo</option>
                        <option value="other">Otro</option>
                    </select>
                </div>
            </div>
            <div class="form-group" style="margin-bottom: 20px;">
                <label>Notas</label>
                <textarea id="notes" rows="2"></textarea>
            </div>
            <button class="btn-primary" onclick="startSession()" style="width: 100%;">🚀 Iniciar</button>
        </div>

        <div class="game-panel hidden" id="gamePanel">
            <div id="gameModeIndicator" class="mode-indicator" style="display: none;"></div>
            <div class="status-bar">
                <div><div style="font-size: 0.9rem; opacity: 0.7;">Paciente</div><div style="font-size: 1.2rem; font-weight: bold;" id="displayUserId">-</div></div>
                <div><div style="font-size: 0.9rem; opacity: 0.7;">Nivel</div><div style="font-size: 1.5rem; font-weight: bold; color: #00d4ff;" id="level">2</div></div>
                <div><div style="font-size: 0.9rem; opacity: 0.7;">Intentos</div><div style="font-size: 1.5rem; font-weight: bold; color: #f59e0b;" id="attempts">2</div></div>
                <div><div style="font-size: 0.9rem; opacity: 0.7;">Span</div><div style="font-size: 1.5rem; font-weight: bold; color: #10b981;" id="span">-</div></div>
            </div>
            <div class="progress-bar"><div class="progress-fill" id="progress" style="width: 0%"></div></div>
            <div class="message info" id="message">Preparando...</div>
            <div class="corsi-board" id="board"></div>
            <div class="controls">
                <button class="btn-primary" id="startBtn" onclick="startLevel()">▶️ Iniciar</button>
                <button class="btn-warning hidden" id="demoBtn" onclick="showDemo()">👁️ Repetir</button>
                <button class="btn-secondary" onclick="abortSession()">⏹️ Terminar</button>
            </div>
        </div>

        <div class="results-panel hidden" id="resultsPanel">
            <h2 style="color: #00d4ff; margin-bottom: 20px;">📊 Completado</h2>
            <div class="metrics-display" id="finalMetrics"></div>
            <div class="download-section">
                <h3>Archivos</h3>
                <p><strong>Modo:</strong> <span id="resultMode"></span></p>
                <div style="display: flex; gap: 10px; justify-content: center; flex-wrap: wrap; margin: 20px 0;">
                    <button class="btn-success" onclick="downloadFile('movements')">📄 Movimientos</button>
                    <button class="btn-success" onclick="downloadFile('summary')">📈 Resumen</button>
                    <button class="btn-success" onclick="downloadFile('combined')">📑 ML CSV</button>
                    <button class="btn-primary" onclick="downloadFile('excel')">📊 Excel</button>
                </div>
                <div id="fileList" class="file-list"><p>Cargando...</p></div>
                <button class="btn-primary" onclick="location.reload()" style="margin-top: 20px;">🔄 Nuevo</button>
            </div>
        </div>
    </div>

    <script>
        const cubePositions = [
            {x: 12, y: 15}, {x: 78, y: 12}, {x: 48, y: 28},
            {x: 22, y: 48}, {x: 68, y: 45}, {x: 88, y: 65},
            {x: 8, y: 72}, {x: 42, y: 78}, {x: 72, y: 82}
        ];

        let testMode = '';
        let sessionId = '';
        let sessionMetadata = {};
        let currentLevel = 2;
        let maxLevel = 9;
        let attemptsLeft = 2;
        let corsiSpan = 0;
        let sequence = [];
        let userSequence = [];
        let isShowingSequence = false;
        let canClick = false;
        let movementsData = [];
        let levelSummaries = [];
        let levelStartTime = 0;
        let sequenceStartTime = 0;
        let lastClickTime = 0;
        let lastClickCoords = {x: 0, y: 0};
        let backtrackingCount = 0;
        let errorCount = 0;
        let currentSelectedCubes = new Set();

        function selectMode(mode) {
            testMode = mode;
            document.querySelectorAll('.mode-card').forEach(c => c.classList.remove('selected'));
            document.getElementById(mode === 'normal' ? 'modeNormal' : 'modeInverse').classList.add('selected');
            setTimeout(() => {
                document.getElementById('modePanel').classList.add('hidden');
                document.getElementById('setupPanel').classList.remove('hidden');
                const modeDisplay = document.getElementById('modeDisplay');
                modeDisplay.style.display = 'inline-block';
                if (mode === 'normal') {
                    modeDisplay.className = 'mode-indicator mode-normal';
                    modeDisplay.textContent = '➡️ DIRECTO';
                } else {
                    modeDisplay.className = 'mode-indicator mode-inverse';
                    modeDisplay.textContent = '🔄 INVERSO';
                }
            }, 300);
        }

        function backToMode() {
            document.getElementById('setupPanel').classList.add('hidden');
            document.getElementById('modePanel').classList.remove('hidden');
            testMode = '';
        }

        function initBoard() {
            const board = document.getElementById('board');
            board.innerHTML = '';
            cubePositions.forEach((pos, idx) => {
                const cube = document.createElement('div');
                cube.className = 'cube';
                cube.id = `cube-${idx}`;
                cube.style.left = `calc(${pos.x}% - 40px)`;
                cube.style.top = `calc(${pos.y}% - 40px)`;
                cube.textContent = idx + 1;
                cube.onclick = (e) => handleCubeClick(idx, e);
                board.appendChild(cube);
            });
        }

        function clearCubeSelection() {
            document.querySelectorAll('.cube').forEach(c => c.classList.remove('selected', 'error', 'active'));
            currentSelectedCubes.clear();
            userSequence = [];
        }

        async function startSession() {
            const userId = document.getElementById('userId').value;
            const age = document.getElementById('age').value;
            if (!userId || !age) { alert('Completa los campos'); return; }

            sessionId = 'SES-' + Date.now() + '-' + Math.random().toString(36).substr(2, 9);
            sessionMetadata = {
                session_id: sessionId, timestamp: new Date().toISOString(),
                user_id: userId, age: parseInt(age),
                gender: document.getElementById('gender').value,
                condition: document.getElementById('condition').value,
                notes: document.getElementById('notes').value,
                test_mode: testMode, final_span: 0, completed: false
            };

            document.getElementById('setupPanel').classList.add('hidden');
            document.getElementById('gamePanel').classList.remove('hidden');
            document.getElementById('displayUserId').textContent = userId;

            const indicator = document.getElementById('gameModeIndicator');
            indicator.style.display = 'inline-block';
            indicator.className = testMode === 'normal' ? 'mode-indicator mode-normal' : 'mode-indicator mode-inverse';
            indicator.textContent = testMode === 'normal' ? '➡️ DIRECTO' : '🔄 INVERSO';

            initBoard();
            logEvent('session_start', null, {test_mode: testMode});
            showMessage('Presiona Iniciar para comenzar', 'info');
        }

        async function startLevel() {
            document.getElementById('startBtn').classList.add('hidden');
            document.getElementById('demoBtn').classList.remove('hidden');
            clearCubeSelection();
            backtrackingCount = 0; errorCount = 0;
            levelStartTime = Date.now();
            
            sequence = [];
            const available = [0,1,2,3,4,5,6,7,8];
            for (let i = 0; i < currentLevel; i++) {
                const randIdx = Math.floor(Math.random() * available.length);
                sequence.push(available[randIdx]);
                available.splice(randIdx, 1);
            }

            updateDisplay();
            showMessage(`Nivel ${currentLevel}: Observa...`, 'info');
            setCubesDisabled(true);
            await sleep(800);
            await showSequence();
        }

        async function showSequence() {
            isShowingSequence = true;
            sequenceStartTime = Date.now();
            logEvent('sequence_start', null, {level: currentLevel, sequence_length: sequence.length, test_mode: testMode});
            
            for (let i = 0; i < sequence.length; i++) {
                const cube = document.getElementById(`cube-${sequence[i]}`);
                cube.classList.add('active');
                playTone(500 + (i * 50), 200);
                await sleep(800);
                cube.classList.remove('active');
                await sleep(300);
            }
            
            isShowingSequence = false;
            canClick = true;
            lastClickTime = Date.now();
            setCubesDisabled(false);
            showMessage(testMode === 'normal' ? 'Repite el MISMO orden' : 'Repite al REVÉS', 'success');
            logEvent('response_phase_start', null, {level: currentLevel, test_mode: testMode});
        }

        async function handleCubeClick(cubeId, event) {
            if (!canClick || isShowingSequence) return;
            const cube = document.getElementById(`cube-${cubeId}`);
            if (currentSelectedCubes.has(cubeId)) {
                logEvent('reclick_attempt', cubeId, {sequence_position: userSequence.length});
                return;
            }

            const now = Date.now();
            const coords = getRelativeCoords(event);
            const reactionTime = userSequence.length === 0 ? now - sequenceStartTime : now - lastClickTime;
            const distance = lastClickCoords.x ? Math.sqrt(Math.pow(coords.x - lastClickCoords.x, 2) + Math.pow(coords.y - lastClickCoords.y, 2)) : 0;
            lastClickTime = now;
            lastClickCoords = coords;

            let expectedCube;
            if (testMode === 'normal') {
                expectedCube = sequence[userSequence.length];
            } else {
                expectedCube = sequence[sequence.length - 1 - userSequence.length];
            }
            
            const isCorrect = cubeId === expectedCube;
            logEvent('cube_click', cubeId, {
                sequence_position: userSequence.length, expected_cube: expectedCube, is_correct: isCorrect,
                reaction_time_ms: reactionTime, distance_px: distance, x_coord: coords.x, y_coord: coords.y,
                level: currentLevel, test_mode: testMode
            });

            if (isCorrect) {
                cube.classList.add('selected');
                currentSelectedCubes.add(cubeId);
                userSequence.push(cubeId);
                playTone(600, 150);
                if (userSequence.length === sequence.length) handleLevelSuccess();
            } else {
                errorCount++;
                cube.classList.add('error');
                playTone(300, 300);
                logEvent('error', cubeId, {error_type: 'commission', expected_cube: expectedCube, clicked_cube: cubeId, test_mode: testMode});
                await sleep(500);
                cube.classList.remove('error');
                handleLevelError();
            }
        }

        function handleLevelSuccess() {
            canClick = false;
            const totalTime = Date.now() - levelStartTime;
            if (currentLevel > corsiSpan) corsiSpan = currentLevel;

            const levelEvents = movementsData.filter(m => m.level === currentLevel && m.event_type === 'cube_click');
            const reactionTimes = levelEvents.map(e => e.reaction_time_ms).filter(t => t);
            const avgReaction = reactionTimes.length ? reactionTimes.reduce((a,b) => a+b, 0) / reactionTimes.length : 0;
            const movements = levelEvents.map(e => e.distance_px / e.reaction_time_ms).filter(v => v && v < 10);
            const avgVelocity = movements.length ? movements.reduce((a,b) => a+b, 0) / movements.length : 0;
            const responseStart = movementsData.find(m => m.event_type === 'response_phase_start' && m.level === currentLevel);
            const firstClick = levelEvents[0];
            const hesitation = (responseStart && firstClick) ? (new Date(firstClick.timestamp) - new Date(responseStart.timestamp)) : 0;

            levelSummaries.push({
                session_id: sessionId, level: currentLevel, success: true, attempts: 3 - attemptsLeft,
                sequence_length: currentLevel, total_time_ms: totalTime, avg_reaction_time_ms: avgReaction,
                first_reaction_time_ms: reactionTimes[0] || 0, hesitation_time_ms: hesitation,
                movement_velocity_avg: avgVelocity, error_count: errorCount, backtracking_count: backtrackingCount,
                test_mode: testMode, timestamp: new Date().toISOString()
            });

            showMessage('¡Correcto! ✅', 'success');
            setTimeout(() => {
                if (currentLevel < maxLevel) {
                    currentLevel++; attemptsLeft = 2; startLevel();
                } else {
                    endSession();
                }
            }, 1200);
        }

        function handleLevelError() {
            attemptsLeft--;
            updateDisplay();
            if (attemptsLeft > 0) {
                showMessage(`Error ❌ Quedan: ${attemptsLeft}`, 'error');
                clearCubeSelection();
                lastClickTime = Date.now();
                setTimeout(() => showSequence(), 1500);
            } else {
                levelSummaries.push({
                    session_id: sessionId, level: currentLevel, success: false, attempts: 2,
                    sequence_length: currentLevel, total_time_ms: Date.now() - levelStartTime,
                    error_count: errorCount, backtracking_count: backtrackingCount,
                    test_mode: testMode, timestamp: new Date().toISOString()
                });
                showMessage(`Nivel ${currentLevel} fallido`, 'error');
                setTimeout(() => endSession(), 1500);
            }
        }

        async function endSession() {
            document.getElementById('gamePanel').classList.add('hidden');
            document.getElementById('resultsPanel').classList.remove('hidden');
            sessionMetadata.final_span = corsiSpan;
            sessionMetadata.completed = true;
            document.getElementById('resultMode').textContent = testMode === 'normal' ? 'Directo' : 'Inverso';

            try {
                const response = await fetch('/api/save_session', {
                    method: 'POST',
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify({metadata: sessionMetadata, movements: movementsData, summaries: levelSummaries})
                });
                const result = await response.json();
                if (result.status === 'ok') await loadFileList();
            } catch (e) { console.error('Error:', e); }

            document.getElementById('finalMetrics').innerHTML = `
                <div class="metric-card"><div class="metric-label">Span</div><div class="metric-value">${corsiSpan}</div></div>
                <div class="metric-card"><div class="metric-label">Nivel</div><div class="metric-value">${currentLevel}</div></div>
                <div class="metric-card"><div class="metric-label">Eventos</div><div class="metric-value">${movementsData.length}</div></div>
                <div class="metric-card"><div class="metric-label">Modo</div><div class="metric-value" style="font-size:1.2rem;">${testMode === 'normal' ? 'Directo' : 'Inverso'}</div></div>
            `;
        }

        function logEvent(eventType, cubeId, data) {
            const baseEvent = {
                session_id: sessionId, timestamp: new Date().toISOString(), event_type: eventType,
                cube_id: cubeId, level: currentLevel, test_mode: testMode,
                sequence_position: null, expected_cube: null, is_correct: null,
                reaction_time_ms: null, distance_px: null, x_coord: null, y_coord: null,
                error_type: null, clicked_cube: null, message: null, sequence_length: null
            };
            movementsData.push({...baseEvent, ...data});
        }

        function updateDisplay() {
            document.getElementById('level').textContent = currentLevel;
            document.getElementById('attempts').textContent = attemptsLeft;
            document.getElementById('span').textContent = corsiSpan || '-';
            document.getElementById('progress').style.width = `${((currentLevel - 2) / 7) * 100}%`;
        }

        function setCubesDisabled(disabled) {
            document.querySelectorAll('.cube').forEach(c => c.classList.toggle('disabled', disabled));
        }

        function getRelativeCoords(event) {
            const rect = document.getElementById('board').getBoundingClientRect();
            return {x: event.clientX - rect.left, y: event.clientY - rect.top};
        }

        async function loadFileList() {
            try {
                const files = await (await fetch('/api/files/' + sessionId)).json();
                document.getElementById('fileList').innerHTML = files.length ? 
                    files.map(f => `<div class="file-item"><span>${f.name}</span><span style="opacity:0.7">${f.size}</span></div>`).join('') : 
                    '<p>No hay archivos</p>';
            } catch (e) { document.getElementById('fileList').innerHTML = '<p style="color:#ef4444">Error</p>'; }
        }

        function downloadFile(type) { window.open(`/api/download/${type}/${sessionId}`, '_blank'); }
        async function abortSession() { if (confirm('¿Terminar?')) endSession(); }
        function sleep(ms) { return new Promise(r => setTimeout(r, ms)); }
        function showMessage(text, type) {
            const msg = document.getElementById('message');
            msg.textContent = text; msg.className = `message ${type}`;
        }
        function playTone(freq, duration) {
            try {
                const audio = new (window.AudioContext || window.webkitAudioContext)();
                const osc = audio.createOscillator(), gain = audio.createGain();
                osc.connect(gain); gain.connect(audio.destination);
                osc.frequency.value = freq;
                gain.gain.setValueAtTime(0.2, audio.currentTime);
                gain.gain.exponentialRampToValueAtTime(0.01, audio.currentTime + duration/1000);
                osc.start(); osc.stop(audio.currentTime + duration/1000);
            } catch(e) {}
        }

        initBoard();
    </script>
</body>
</html>
"""

@app.route('/')
def index():
    return render_template_string(HTML_TEMPLATE)

@app.route('/api/save_session', methods=['POST'])
def save_session():
    try:
        data = request.json
        session_id = data['metadata']['session_id']
        
        print(f"💾 Guardando: {session_id}")
        
        # 1. Metadata
        session_file = os.path.join(DATA_DIR, 'sessions', f'{session_id}_metadata.csv')
        with open(session_file, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=data['metadata'].keys())
            writer.writeheader()
            writer.writerow(data['metadata'])
        
        # 2. Movimientos
        if data['movements']:
            all_fields = set()
            for mov in data['movements']:
                all_fields.update(mov.keys())
            movements_file = os.path.join(DATA_DIR, 'movements', f'{session_id}_movements.csv')
            with open(movements_file, 'w', newline='', encoding='utf-8') as f:
                writer = csv.DictWriter(f, fieldnames=sorted(list(all_fields)))
                writer.writeheader()
                writer.writerows(data['movements'])
        
        # 3. Summary
        if data['summaries']:
            summary_file = os.path.join(DATA_DIR, 'movements', f'{session_id}_summary.csv')
            with open(summary_file, 'w', newline='', encoding='utf-8') as f:
                writer = csv.DictWriter(f, fieldnames=data['summaries'][0].keys())
                writer.writeheader()
                writer.writerows(data['summaries'])
        
        # 4. ML Dataset
        ml_file = create_ml_dataset(session_id, data)
        
        # 5. Excel
        excel_file = create_simple_excel(session_id, data)
        
        return jsonify({'status': 'ok', 'files_saved': 5, 'data_path': DATA_DIR, 'session_id': session_id})
        
    except Exception as e:
        print(f"❌ Error: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({'status': 'error', 'message': str(e)}), 500

def create_ml_dataset(session_id, data):
    """CORREGIDO: successes es int, no lista"""
    features = {
        'session_id': session_id,
        'user_id': data['metadata']['user_id'],
        'age': data['metadata']['age'],
        'gender': data['metadata']['gender'],
        'condition': data['metadata']['condition'],
        'test_mode': data['metadata'].get('test_mode', 'normal'),
        'final_span': data['metadata']['final_span'],
        'total_events': len(data['movements']),
        'total_levels': len(data['summaries']),
    }
    
    if data['movements']:
        clicks = [m for m in data['movements'] if m.get('event_type') == 'cube_click']
        if clicks:
            reaction_times = [m.get('reaction_time_ms', 0) for m in clicks if m.get('reaction_time_ms')]
            distances = [m.get('distance_px', 0) for m in clicks if m.get('distance_px')]
            error_count = len([m for m in clicks if m.get('is_correct') == False])
            
            if reaction_times:
                features['mean_reaction_time'] = sum(reaction_times) / len(reaction_times)
                features['std_reaction_time'] = pd.Series(reaction_times).std() if len(reaction_times) > 1 else 0
            if distances:
                features['mean_distance'] = sum(distances) / len(distances)
            
            features['total_errors'] = error_count
            features['error_rate'] = error_count / len(clicks) if clicks else 0
            
            velocities = []
            for m in clicks:
                rt = m.get('reaction_time_ms', 0)
                dist = m.get('distance_px', 0)
                if rt and rt > 0 and dist:
                    velocities.append(dist / rt)
            
            if velocities:
                features['mean_velocity'] = sum(velocities) / len(velocities)
                features['max_velocity'] = max(velocities)
    
    if data['summaries']:
        hesitations = [s.get('hesitation_time_ms', 0) for s in data['summaries'] if s.get('hesitation_time_ms')]
        backtracks = sum([s.get('backtracking_count', 0) for s in data['summaries']])
        # CORRECCIÓN: successes es un CONTADOR (int), no una lista
        success_count = len([s for s in data['summaries'] if s.get('success')])
        times = [s.get('total_time_ms', 0) for s in data['summaries'] if s.get('total_time_ms')]
        
        if hesitations:
            features['mean_hesitation_time'] = sum(hesitations) / len(hesitations)
        
        features['total_backtracks'] = backtracks
        # CORRECCIÓN: usar success_count directamente, no len(successes)
        features['success_rate'] = success_count / len(data['summaries']) if data['summaries'] else 0
        
        if times:
            features['mean_level_time'] = sum(times) / len(times)
    
    ml_file = os.path.join(DATA_DIR, 'exports', f'{session_id}_ml_ready.csv')
    pd.DataFrame([features]).to_csv(ml_file, index=False)
    return ml_file

def create_simple_excel(session_id, data):
    """Excel simple sin celdas mergeadas"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    
    excel_path = os.path.join(DATA_DIR, 'exports', f'{session_id}_completo.xlsx')
    wb = Workbook()
    
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    title_font = Font(size=14, bold=True, color="366092")
    section_font = Font(bold=True, size=11, color="366092")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                    top=Side(style='thin'), bottom=Side(style='thin'))
    
    # HOJA 1: RESUMEN
    ws1 = wb.active
    ws1.title = "Resumen"
    ws1['A1'] = "PRUEBA DE CORSI - RESUMEN"
    ws1['A1'].font = title_font
    ws1['A3'] = "Información General"
    ws1['A3'].font = section_font
    
    row = 4
    for label, value in [
        ["ID Sesión:", data['metadata']['session_id']],
        ["Paciente:", data['metadata']['user_id']],
        ["Edad:", data['metadata']['age']],
        ["Género:", data['metadata']['gender']],
        ["Condición:", data['metadata']['condition']],
        ["Modo:", "Directo" if data['metadata'].get('test_mode') == 'normal' else "Inverso"],
        ["Span Final:", data['metadata']['final_span']],
        ["Fecha:", data['metadata']['timestamp'][:10]],
        ["Hora:", data['metadata']['timestamp'][11:19]],
    ]:
        ws1.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws1.cell(row=row, column=2, value=value)
        row += 1
    
    row += 2
    ws1.cell(row=row, column=1, value="Métricas").font = section_font
    row += 1
    
    if data['movements']:
        clicks = [m for m in data['movements'] if m.get('event_type') == 'cube_click']
        error_list = [m for m in clicks if m.get('is_correct') == False]
        reaction_times = [m.get('reaction_time_ms', 0) for m in clicks if m.get('reaction_time_ms')]
        
        for label, value in [
            ["Total clicks:", len(clicks)],
            ["Errores:", len(error_list)],
            ["Tasa error:", f"{(len(error_list)/len(clicks)*100):.1f}%" if clicks else "0%"],
            ["React medio:", f"{(sum(reaction_times)/len(reaction_times)):.0f} ms" if reaction_times else "N/A"],
        ]:
            ws1.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws1.cell(row=row, column=2, value=value)
            row += 1
    
    ws1.column_dimensions['A'].width = 20
    ws1.column_dimensions['B'].width = 30
    
    # HOJA 2: MOVIMIENTOS
    ws2 = wb.create_sheet("Movimientos")
    if data['movements']:
        headers = ['Timestamp', 'Evento', 'Nivel', 'Cubo', 'Pos', 'Esperado', 'Correcto', 'React(ms)', 'Dist(px)', 'X', 'Y']
        for col, header in enumerate(headers, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
        
        for row_idx, mov in enumerate(data['movements'], 2):
            ws2.cell(row=row_idx, column=1, value=str(mov.get('timestamp', ''))[:19])
            ws2.cell(row=row_idx, column=2, value=mov.get('event_type', ''))
            ws2.cell(row=row_idx, column=3, value=mov.get('level', ''))
            ws2.cell(row=row_idx, column=4, value=mov.get('cube_id', ''))
            ws2.cell(row=row_idx, column=5, value=mov.get('sequence_position', ''))
            ws2.cell(row=row_idx, column=6, value=mov.get('expected_cube', ''))
            
            correct = mov.get('is_correct')
            cell = ws2.cell(row=row_idx, column=7, value='Sí' if correct else ('No' if correct == False else ''))
            if correct == True:
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            elif correct == False:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            
            ws2.cell(row=row_idx, column=8, value=mov.get('reaction_time_ms', ''))
            ws2.cell(row=row_idx, column=9, value=mov.get('distance_px', ''))
            ws2.cell(row=row_idx, column=10, value=mov.get('x_coord', ''))
            ws2.cell(row=row_idx, column=11, value=mov.get('y_coord', ''))
            
            for col in range(1, 12):
                ws2.cell(row=row_idx, column=col).border = border
        
        for col in range(1, 12):
            ws2.column_dimensions[get_column_letter(col)].width = 12
    
    # HOJA 3: NIVELES
    ws3 = wb.create_sheet("Niveles")
    if data['summaries']:
        headers = ['Nivel', 'Éxito', 'Intentos', 'Tiempo(ms)', 'ReactProm', 'React1er', 'Hesitación', 'Velocidad', 'Errores', 'Backtracks']
        for col, header in enumerate(headers, 1):
            cell = ws3.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
        
        for row_idx, summary in enumerate(data['summaries'], 2):
            ws3.cell(row=row_idx, column=1, value=summary.get('level'))
            
            success = summary.get('success')
            cell = ws3.cell(row=row_idx, column=2, value='Sí' if success else 'No')
            if success:
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            else:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            
            ws3.cell(row=row_idx, column=3, value=summary.get('attempts'))
            ws3.cell(row=row_idx, column=4, value=summary.get('total_time_ms'))
            
            avg_rt = summary.get('avg_reaction_time_ms')
            ws3.cell(row=row_idx, column=5, value=round(avg_rt, 2) if avg_rt else '')
            
            first_rt = summary.get('first_reaction_time_ms')
            ws3.cell(row=row_idx, column=6, value=round(first_rt, 2) if first_rt else '')
            
            hes = summary.get('hesitation_time_ms')
            ws3.cell(row=row_idx, column=7, value=round(hes, 2) if hes else '')
            
            vel = summary.get('movement_velocity_avg')
            ws3.cell(row=row_idx, column=8, value=round(vel, 4) if vel else '')
            
            ws3.cell(row=row_idx, column=9, value=summary.get('error_count', 0))
            ws3.cell(row=row_idx, column=10, value=summary.get('backtracking_count', 0))
            
            for col in range(1, 11):
                ws3.cell(row=row_idx, column=col).border = border
        
        for col in range(1, 11):
            ws3.column_dimensions[get_column_letter(col)].width = 12
    
    # HOJA 4: FEATURES
    ws4 = wb.create_sheet("Features ML")
    row = 1
    ws4.cell(row=row, column=1, value="FEATURES PARA ML").font = title_font
    row += 2
    
    features_list = [("Categoría", "Feature", "Valor", "Descripción")]
    features_list.append(("", "", "", ""))
    features_list.append(("Demográficos", "Edad", data['metadata']['age'], "Edad"))
    features_list.append(("Demográficos", "Género", data['metadata']['gender'], "M/F/O"))
    features_list.append(("Demográficos", "Condición", data['metadata']['condition'], "Grupo"))
    features_list.append(("", "", "", ""))
    features_list.append(("Rendimiento", "Span", data['metadata']['final_span'], "Máximo nivel"))
    features_list.append(("Rendimiento", "Modo", "Directo" if data['metadata'].get('test_mode') == 'normal' else "Inverso", "Tipo"))
    
    if data['movements']:
        clicks = [m for m in data['movements'] if m.get('event_type') == 'cube_click']
        error_list = [m for m in clicks if m.get('is_correct') == False]
        reaction_times = [m.get('reaction_time_ms', 0) for m in clicks if m.get('reaction_time_ms')]
        
        features_list.append(("Rendimiento", "Clicks", len(clicks), "Total interacciones"))
        features_list.append(("Rendimiento", "Errores", len(error_list), "Total errores"))
        features_list.append(("Rendimiento", "Tasa Error", f"{(len(error_list)/len(clicks)*100):.2f}%" if clicks else "0%", "Porcentaje"))
        features_list.append(("", "", "", ""))
        features_list.append(("Temporales", "React Medio", f"{(sum(reaction_times)/len(reaction_times)):.2f} ms" if reaction_times else "N/A", "Promedio"))
        
        velocities = []
        for m in clicks:
            rt = m.get('reaction_time_ms', 0)
            dist = m.get('distance_px', 0)
            if rt and rt > 0 and dist:
                velocities.append(dist / rt)
        
        if velocities:
            features_list.append(("Espaciales", "Velocidad Media", f"{(sum(velocities)/len(velocities)):.4f}", "px/ms"))
    
    if data['summaries']:
        hesitations = [s.get('hesitation_time_ms', 0) for s in data['summaries'] if s.get('hesitation_time_ms')]
        backtracks = sum([s.get('backtracking_count', 0) for s in data['summaries']])
        success_count = len([s for s in data['summaries'] if s.get('success')])
        
        features_list.append(("", "", "", ""))
        features_list.append(("Comportamentales", "Hesitación", f"{(sum(hesitations)/len(hesitations)):.2f} ms" if hesitations else "N/A", "Tiempo inicial"))
        features_list.append(("Comportamentales", "Backtracks", backtracks, "Repeticiones"))
        features_list.append(("Comportamentales", "Exitosos", success_count, "Niveles completados"))
    
    for r_idx, (cat, feat, val, desc) in enumerate(features_list, row):
        ws4.cell(row=r_idx, column=1, value=cat)
        ws4.cell(row=r_idx, column=2, value=feat)
        ws4.cell(row=r_idx, column=3, value=val)
        ws4.cell(row=r_idx, column=4, value=desc)
        
        if r_idx == row:
            for c in range(1, 5):
                cell = ws4.cell(row=r_idx, column=c)
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
    
    ws4.column_dimensions['A'].width = 16
    ws4.column_dimensions['B'].width = 16
    ws4.column_dimensions['C'].width = 16
    ws4.column_dimensions['D'].width = 30
    
    wb.save(excel_path)
    print(f"✅ Excel: {excel_path}")
    return excel_path

@app.route('/api/files/<session_id>')
def list_files(session_id):
    files = []
    folders = {
        'sessions': os.path.join(DATA_DIR, 'sessions'),
        'movements': os.path.join(DATA_DIR, 'movements'),
        'exports': os.path.join(DATA_DIR, 'exports')
    }
    for folder_name, folder_path in folders.items():
        if os.path.exists(folder_path):
            for f in os.listdir(folder_path):
                if session_id in f:
                    path = os.path.join(folder_path, f)
                    size = os.path.getsize(path)
                    files.append({'name': f, 'size': f'{size/1024:.1f} KB', 'path': path})
    return jsonify(files)

@app.route('/api/download/<type>/<session_id>')
def download_file(type, session_id):
    try:
        if type == 'movements':
            filename = f'{session_id}_movements.csv'
            filepath = os.path.join(DATA_DIR, 'movements', filename)
            mimetype = 'text/csv'
        elif type == 'summary':
            filename = f'{session_id}_summary.csv'
            filepath = os.path.join(DATA_DIR, 'movements', filename)
            mimetype = 'text/csv'
        elif type == 'combined':
            filename = f'{session_id}_ml_ready.csv'
            filepath = os.path.join(DATA_DIR, 'exports', filename)
            mimetype = 'text/csv'
        elif type == 'excel':
            filename = f'{session_id}_completo.xlsx'
            filepath = os.path.join(DATA_DIR, 'exports', filename)
            mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        else:
            return 'Tipo no válido', 400
        
        if os.path.exists(filepath):
            try:
                return send_file(filepath, as_attachment=True, attachment_filename=filename, mimetype=mimetype)
            except TypeError:
                from flask import Response
                with open(filepath, 'rb') as f:
                    return Response(f.read(), mimetype=mimetype, headers={'Content-Disposition': f'attachment; filename={filename}'})
        return f'No encontrado: {filename}', 404
    except Exception as e:
        return str(e), 500

if __name__ == '__main__':
    print("🧠 Corsi ML v3.3 - http://localhost:8080")
    app.run(debug=False, host='0.0.0.0', port=8080, threaded=True)