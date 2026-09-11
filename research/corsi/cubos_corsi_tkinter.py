import tkinter as tk
from tkinter import ttk, messagebox
import random
import time
import os
import sys
import csv
import statistics
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import threading

# Configuración rutas
if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DATA_DIR = os.path.join(BASE_DIR, 'Corsi_Resultados')
os.makedirs(DATA_DIR, exist_ok=True)


class CorsiApp:
    def __init__(self, root):
        self.root = root
        self.root.title("CORSI - Evaluación de Memoria Visoespacial")
        self.root.geometry("1400x900")
        self.root.configure(bg='#1a1c2e')
        self.root.minsize(1200, 800)
        
        # Paleta morada/glassmorphism
        self.colors = {
            'bg': '#1a1c2e',
            'card': '#2d2a4a',
            'input_bg': '#0f172a',
            'accent': '#00d4ff',
            'success': '#10b981',
            'warning': '#f59e0b',
            'error': '#ef4444',
            'text': '#f8fafc',
            'muted': '#94a3b8',
            'border': '#4a5568',
            'board': '#0a0f1a',
            'cube': '#1e293b'
        }
        
        # Variables
        self.test_mode = ''
        self.session_id = ''
        self.metadata = {}
        self.current_level = 2
        self.attempts_left = 2
        self.corsi_span = 0
        self.sequence = []
        self.user_sequence = []
        self.is_showing = False
        self.can_click = False
        self.movements = []
        self.summaries = []
        self.cubes = []
        self.active_timers = []
        self.level_in_progress = False
        self.level_start_time = 0
        self.sequence_start_time = 0
        self.last_click_time = 0
        self.first_level_started = False
        self.excel_path = None
        self.sound_initialized = False  # Control de inicialización de sonido
        
        # Posiciones cubos
        self.positions = [
            (20, 18), (80, 15), (50, 30),
            (25, 50), (75, 45), (85, 70),
            (15, 75), (48, 82), (78, 85)
        ]
        
        # Inicializar sistema de audio en hilo separado para evitar bloqueos
        self.init_audio()
        
        self.show_welcome()
    
    def init_audio(self):
        """Inicializar sistema de audio de forma segura"""
        def warmup():
            try:
                import winsound
                # Sonido de prueba muy corto e imperceptible para "despertar" el sistema
                winsound.Beep(100, 50)
                self.sound_initialized = True
            except:
                pass
        
        # Ejecutar en hilo separado para no bloquear la UI
        threading.Thread(target=warmup, daemon=True).start()
    
    def clear(self):
        """Limpiar pantalla y cancelar timers"""
        for timer_id in self.active_timers:
            try:
                self.root.after_cancel(timer_id)
            except:
                pass
        self.active_timers = []
        
        for widget in self.root.winfo_children():
            widget.destroy()
    
    def safe_after(self, ms, func):
        """Timer seguro"""
        timer_id = self.root.after(ms, func)
        self.active_timers.append(timer_id)
        return timer_id
    
    def create_rounded_button(self, parent, text, command, bg_color, fg_color, 
                             width=200, height=45, font_size=12, bold=True):
        """Botón redondeado con Canvas"""
        canvas = tk.Canvas(parent, width=width, height=height, 
                          bg=parent['bg'], highlightthickness=0, cursor='hand2')
        
        r = 10
        points = [r, 0, width-r, 0, width, 0, width, r, width, height-r, 
                 width, height, width-r, height, r, height, 0, height, 
                 0, height-r, 0, r, 0, 0]
        
        rect_id = canvas.create_polygon(points, smooth=True, fill=bg_color, outline='')
        
        weight = 'bold' if bold else 'normal'
        text_id = canvas.create_text(width//2, height//2, text=text, 
                                    fill=fg_color, font=('Segoe UI', font_size, weight))
        
        def on_enter(e):
            canvas.itemconfig(rect_id, fill=self.lighten_color(bg_color, 15))
        def on_leave(e):
            canvas.itemconfig(rect_id, fill=bg_color)
        def on_click(e):
            if command:
                command()
        
        canvas.bind('<Enter>', on_enter)
        canvas.bind('<Leave>', on_leave)
        canvas.bind('<Button-1>', on_click)
        
        return canvas
    
    def lighten_color(self, hex_color, percent):
        """Aclarar color"""
        hex_color = hex_color.lstrip('#')
        rgb = tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
        new_rgb = tuple(min(255, int(c + (255-c) * percent/100)) for c in rgb)
        return '#{:02x}{:02x}{:02x}'.format(*new_rgb)
    
    def show_welcome(self):
        """Pantalla de bienvenida con historial"""
        self.clear()
        
        # Header
        header = tk.Frame(self.root, bg='#252545', height=120)
        header.pack(fill=tk.X)
        header.pack_propagate(False)
        
        header_content = tk.Frame(header, bg='#252545')
        header_content.place(relx=0.5, rely=0.5, anchor='center')
        
        tk.Label(header_content, text="🧠 CORSI", 
                font=('Segoe UI', 32, 'bold'), fg='#60a5fa', bg='#252545').pack()
        
        tk.Label(header_content, text="Evaluación de Memoria Visoespacial",
                font=('Segoe UI', 12), fg='#94a3b8', bg='#252545').pack()
        
        # Badge
        badge = tk.Frame(header_content, bg='#10b981', padx=10, pady=2)
        badge.pack(pady=(8, 0))
        tk.Label(badge, text="Sistema ML-Ready", font=('Segoe UI', 9, 'bold'),
                fg='white', bg='#10b981').pack()
        
        # Contenido principal
        main = tk.Frame(self.root, bg=self.colors['bg'])
        main.pack(fill=tk.BOTH, expand=True, padx=80, pady=30)
        
        card = tk.Frame(main, bg='#2d2a4a', highlightbackground='#4a5568', 
                       highlightthickness=1, padx=40, pady=30)
        card.pack(fill=tk.BOTH, expand=True)
        
        # Título
        tk.Label(card, text="Selecciona el Modo de Prueba", 
                font=('Segoe UI', 16, 'bold'), fg=self.colors['accent'], 
                bg='#2d2a4a').pack(pady=(0, 25))
        
        # Frame para modos
        modes_frame = tk.Frame(card, bg='#2d2a4a')
        modes_frame.pack()
        
        self.create_mode_card(modes_frame, 'normal', '➡️', 'Corsi Directo',
                             'Repite la secuencia en el mismo orden',
                             self.colors['accent'], 0)
        
        self.create_mode_card(modes_frame, 'inverse', '🔄', 'Corsi Inverso',
                             'Repite la secuencia en orden inverso',
                             self.colors['warning'], 1)
        
        # Botón de historial (abajo)
        tk.Frame(card, bg='#2d2a4a', height=20).pack()  # Espaciador
        
        history_btn = self.create_rounded_button(card, "📂 Ver Historial de Pruebas", 
                                                self.show_history,
                                                '#4a5568', 'white', 
                                                width=300, height=45, font_size=12)
        history_btn.pack(pady=(20, 0))
    
    def create_mode_card(self, parent, mode, icon, title, desc, color, col):
        """Tarjeta de modo"""
        card = tk.Frame(parent, bg='#3d3a5c', width=280, height=200,
                       highlightbackground='#4a5568', highlightthickness=1)
        card.grid(row=0, column=col, padx=12, pady=5)
        card.grid_propagate(False)
        card.configure(cursor='hand2')
        
        def on_click(event=None):
            self.select_mode(mode)
        
        def on_enter(e):
            card.configure(bg='#4d4a6c')
            inner.configure(bg='#4d4a6c')
            for w in inner.winfo_children():
                w.configure(bg='#4d4a6c')
        
        def on_leave(e):
            card.configure(bg='#3d3a5c')
            inner.configure(bg='#3d3a5c')
            for w in inner.winfo_children():
                w.configure(bg='#3d3a5c')
        
        card.bind('<Enter>', on_enter)
        card.bind('<Leave>', on_leave)
        card.bind('<Button-1>', on_click)
        
        inner = tk.Frame(card, bg='#3d3a5c')
        inner.place(relx=0.5, rely=0.5, anchor='center')
        
        lbl_icon = tk.Label(inner, text=icon, font=('Segoe UI', 48),
                           bg='#3d3a5c', fg=color, cursor='hand2')
        lbl_icon.pack()
        lbl_icon.bind('<Button-1>', on_click)
        
        lbl_title = tk.Label(inner, text=title, font=('Segoe UI', 16, 'bold'),
                            bg='#3d3a5c', fg=color, cursor='hand2')
        lbl_title.pack(pady=(10, 5))
        lbl_title.bind('<Button-1>', on_click)
        
        lbl_desc = tk.Label(inner, text=desc, font=('Segoe UI', 10),
                           bg='#3d3a5c', fg='#cbd5e1', cursor='hand2', wraplength=240)
        lbl_desc.pack()
        lbl_desc.bind('<Button-1>', on_click)
    
    def select_mode(self, mode):
        self.test_mode = mode
        self.show_config()
    
    def show_config(self):
        """Pantalla de configuración con Nombre y sin Notas"""
        self.clear()
        
        # Header
        header = tk.Frame(self.root, bg='#252545', height=100)
        header.pack(fill=tk.X)
        header.pack_propagate(False)
        
        header_content = tk.Frame(header, bg='#252545')
        header_content.place(relx=0.5, rely=0.5, anchor='center')
        
        tk.Label(header_content, text="🧠 CORSI", 
                font=('Segoe UI', 24, 'bold'), fg='#60a5fa', bg='#252545').pack()
        tk.Label(header_content, text="Evaluación de Memoria Visoespacial",
                font=('Segoe UI', 10), fg='#94a3b8', bg='#252545').pack()
        
        # Contenido
        main = tk.Frame(self.root, bg=self.colors['bg'])
        main.pack(fill=tk.BOTH, expand=True, padx=80, pady=20)
        
        card = tk.Frame(main, bg='#2d2a4a', highlightbackground='#4a5568',
                       highlightthickness=1, padx=40, pady=30)
        card.pack(fill=tk.BOTH, expand=True)
        
        # Header card
        card_header = tk.Frame(card, bg='#2d2a4a')
        card_header.pack(fill=tk.X, pady=(0, 20))
        
        tk.Label(card_header, text="Configuración del Participante", 
                font=('Segoe UI', 18, 'bold'), fg=self.colors['accent'], 
                bg='#2d2a4a').pack(side=tk.LEFT)
        
        btn_change = self.create_rounded_button(card_header, "← CAMBIAR MODO", 
                                               self.show_welcome,
                                               '#4a5568', 'white', width=160, height=35, 
                                               font_size=10)
        btn_change.pack(side=tk.RIGHT)
        
        # Indicador modo
        mode_color = self.colors['accent'] if self.test_mode == 'normal' else self.colors['warning']
        mode_text = "➡️ DIRECTO" if self.test_mode == 'normal' else "🔄 INVERSO"
        
        mode_badge = tk.Frame(card, bg='#1a1c2e', highlightbackground=mode_color,
                             highlightthickness=2, padx=15, pady=5)
        mode_badge.pack(anchor='w', pady=(0, 20))
        tk.Label(mode_badge, text=mode_text, font=('Segoe UI', 10, 'bold'),
                fg=mode_color, bg='#1a1c2e').pack()
        
        # Formulario - SIN NOTAS
        form = tk.Frame(card, bg='#2d2a4a')
        form.pack(fill=tk.X)
        
        # Fila 1: ID y Nombre
        self.entry_id = self.create_input(form, "ID Participante *", 0, 0, 20, "Ej: PAC-001")
        self.entry_name = self.create_input(form, "Nombre del Participante", 0, 1, 25, "Ej: Juan Pérez")
        
        # Fila 2: Edad, Género, Condición
        self.entry_age = self.create_input(form, "Edad *", 1, 0, 12)
        self.combo_gender = self.create_combo(form, "Género", 1, 1, 
                                              ['Masculino', 'Femenino', 'Otro'])
        self.combo_condition = self.create_combo(form, "Condición", 1, 2,
                                                 ['Control', 'TDAH', 'Deterioro Cognitivo', 
                                                  'Alzheimer', 'Traumatismo', 'Otro'])
        
        # Botón iniciar
        btn_frame = tk.Frame(card, bg='#2d2a4a')
        btn_frame.pack(fill=tk.X, pady=(30, 0))
        
        btn_start = self.create_rounded_button(btn_frame, "🚀  INICIAR PRUEBA", 
                                              self.validate_start,
                                              self.colors['accent'], '#0f172a',
                                              width=400, height=50, font_size=14)
        btn_start.pack()
    
    def create_input(self, parent, label, row, col, width, placeholder=""):
        frame = tk.Frame(parent, bg='#2d2a4a')
        frame.grid(row=row, column=col, padx=10, pady=10, sticky='ew')
        
        tk.Label(frame, text=label, font=('Segoe UI', 10, 'bold'),
                fg=self.colors['accent'], bg='#2d2a4a').pack(anchor='w')
        
        entry = tk.Entry(frame, font=('Segoe UI', 11), width=width,
                        bg=self.colors['input_bg'], fg='white',
                        insertbackground='white', relief=tk.FLAT,
                        highlightthickness=1, highlightbackground='#4a5568')
        entry.pack(fill=tk.X, pady=(5, 0), ipady=8)
        
        if placeholder:
            entry.insert(0, placeholder)
            entry.bind('<FocusIn>', lambda e: entry.delete(0, tk.END) if entry.get() == placeholder else None)
        
        return entry
    
    def create_combo(self, parent, label, row, col, values):
        frame = tk.Frame(parent, bg='#2d2a4a')
        frame.grid(row=row, column=col, padx=10, pady=10, sticky='ew')
        
        tk.Label(frame, text=label, font=('Segoe UI', 10, 'bold'),
                fg=self.colors['accent'], bg='#2d2a4a').pack(anchor='w')
        
        style = ttk.Style()
        style.theme_use('clam')
        style.configure('Custom.TCombobox', 
                       fieldbackground=self.colors['input_bg'],
                       background=self.colors['input_bg'],
                       foreground='white',
                       arrowcolor='white')
        
        combo = ttk.Combobox(frame, values=values, state='readonly', 
                            font=('Segoe UI', 11), style='Custom.TCombobox')
        combo.set(values[0])
        combo.pack(fill=tk.X, pady=(5, 0))
        
        return combo
    
    def validate_start(self):
        user_id = self.entry_id.get().strip()
        user_name = self.entry_name.get().strip()
        age = self.entry_age.get().strip()
        
        if user_id == "Ej: PAC-001":
            user_id = ""
        if user_name == "Ej: Juan Pérez":
            user_name = ""
        
        if not user_id or not age:
            messagebox.showerror("Error", "Completa ID y Edad (obligatorios)")
            return
        
        try:
            age = int(age)
            if not (5 <= age <= 100):
                raise ValueError
        except:
            messagebox.showerror("Error", "Edad inválida (5-100)")
            return
        
        self.session_id = f"CORSI_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        self.metadata = {
            'session_id': self.session_id,
            'timestamp': datetime.now().isoformat(),
            'user_id': user_id,
            'user_name': user_name if user_name else "No especificado",
            'age': age,
            'gender': self.combo_gender.get(),
            'condition': self.combo_condition.get(),
            'test_mode': self.test_mode,
            'final_span': 0,
            'total_levels_attempted': 0,
            'total_levels_success': 0
        }
        
        # Resetear variables de control
        self.first_level_started = False
        self.current_level = 2
        self.attempts_left = 2
        self.corsi_span = 0
        self.movements = []
        self.summaries = []
        self.excel_path = None
        
        self.show_game()
    
    def show_game(self):
        """Pantalla de juego"""
        self.clear()
        
        # Card principal
        main_card = tk.Frame(self.root, bg='#2d2a4a', highlightbackground='#4a5568',
                            highlightthickness=1)
        main_card.pack(fill=tk.BOTH, expand=True, padx=40, pady=20)
        
        # Header
        header = tk.Frame(main_card, bg='#2d2a4a', padx=20, pady=15)
        header.pack(fill=tk.X)
        
        mode_color = self.colors['accent'] if self.test_mode == 'normal' else self.colors['warning']
        mode_text = "➡️ DIRECTO" if self.test_mode == 'normal' else "🔄 INVERSO"
        
        mode_badge = tk.Frame(header, bg='#1a1c2e', highlightbackground=mode_color,
                             highlightthickness=2, padx=12, pady=4)
        mode_badge.pack(side=tk.LEFT)
        tk.Label(mode_badge, text=mode_text, font=('Segoe UI', 9, 'bold'),
                fg=mode_color, bg='#1a1c2e').pack()
        
        # Métricas
        metrics_bar = tk.Frame(main_card, bg='#1a1c2e', padx=20, pady=12)
        metrics_bar.pack(fill=tk.X, padx=20, pady=(0, 10))
        
        metrics_grid = tk.Frame(metrics_bar, bg='#1a1c2e')
        metrics_grid.pack(fill=tk.X)
        
        tk.Label(metrics_grid, text="Paciente", font=('Segoe UI', 9),
                fg='#64748b', bg='#1a1c2e').grid(row=0, column=0, sticky='w')
        tk.Label(metrics_grid, text=self.metadata['user_name'][:15], font=('Segoe UI', 12, 'bold'),
                fg='white', bg='#1a1c2e').grid(row=1, column=0, sticky='w')
        
        tk.Label(metrics_grid, text="ID", font=('Segoe UI', 9),
                fg='#64748b', bg='#1a1c2e').grid(row=0, column=1, padx=(60, 0), sticky='w')
        tk.Label(metrics_grid, text=self.metadata['user_id'], font=('Segoe UI', 10),
                fg='#94a3b8', bg='#1a1c2e').grid(row=1, column=1, padx=(60, 0), sticky='w')
        
        tk.Label(metrics_grid, text="Nivel", font=('Segoe UI', 9),
                fg='#64748b', bg='#1a1c2e').grid(row=0, column=2, padx=(60, 0), sticky='w')
        self.lbl_level = tk.Label(metrics_grid, text=str(self.current_level), 
                                 font=('Segoe UI', 16, 'bold'),
                                 fg=self.colors['accent'], bg='#1a1c2e')
        self.lbl_level.grid(row=1, column=2, padx=(60, 0), sticky='w')
        
        tk.Label(metrics_grid, text="Intentos", font=('Segoe UI', 9),
                fg='#64748b', bg='#1a1c2e').grid(row=0, column=3, padx=(60, 0), sticky='w')
        self.lbl_attempts = tk.Label(metrics_grid, text=str(self.attempts_left), 
                                    font=('Segoe UI', 16, 'bold'),
                                    fg=self.colors['warning'], bg='#1a1c2e')
        self.lbl_attempts.grid(row=1, column=3, padx=(60, 0), sticky='w')
        
        tk.Label(metrics_grid, text="Span", font=('Segoe UI', 9),
                fg='#64748b', bg='#1a1c2e').grid(row=0, column=4, padx=(60, 0), sticky='w')
        self.lbl_span = tk.Label(metrics_grid, text="-", font=('Segoe UI', 16, 'bold'),
                                fg=self.colors['success'], bg='#1a1c2e')
        self.lbl_span.grid(row=1, column=4, padx=(60, 0), sticky='w')
        
        # Mensaje
        self.status_frame = tk.Frame(main_card, bg='#3d3a5c', padx=20, pady=12)
        self.status_frame.pack(fill=tk.X, padx=20, pady=10)
        
        self.lbl_message = tk.Label(self.status_frame, 
                                   text="Presiona INICIAR para comenzar el nivel",
                                   font=('Segoe UI', 12), fg='white', bg='#3d3a5c')
        self.lbl_message.pack()
        
        # Tablero
        board_container = tk.Frame(main_card, bg=self.colors['board'], 
                                  highlightbackground=self.colors['accent'], 
                                  highlightthickness=2)
        board_container.pack(padx=20, pady=10, fill=tk.BOTH, expand=True)
        
        self.board_canvas = tk.Canvas(board_container, bg=self.colors['board'],
                                     highlightthickness=0)
        self.board_canvas.pack(fill=tk.BOTH, expand=True)
        
        self.root.after(100, self.draw_cubes)
        
        # Controles
        controls = tk.Frame(main_card, bg='#2d2a4a', pady=20)
        controls.pack(fill=tk.X)
        
        self.btn_start = self.create_rounded_button(controls, "▶️  INICIAR", 
                                                   self.start_level,
                                                   self.colors['accent'], '#0f172a',
                                                   width=180, height=45, font_size=12)
        self.btn_start.pack(side=tk.LEFT, padx=(20, 10))
        
        self.btn_repeat = self.create_rounded_button(controls, "👁️  REPETIR", 
                                                    self.repeat_sequence,
                                                    '#4a5568', 'white',
                                                    width=160, height=45, font_size=11)
        self.btn_repeat.pack(side=tk.LEFT, padx=10)
        
        btn_stop = self.create_rounded_button(controls, "⏹️  TERMINAR", 
                                             self.confirm_stop,
                                             '#ef4444', 'white',
                                             width=160, height=45, font_size=11)
        btn_stop.pack(side=tk.RIGHT, padx=20)
        
        self.log_event('session_start', None, {'mode': self.test_mode})
        self.level_in_progress = False
    
    def draw_cubes(self):
        """Dibujar cubos"""
        width = self.board_canvas.winfo_width()
        height = self.board_canvas.winfo_height()
        
        if width < 100 or height < 100:
            self.root.after(100, self.draw_cubes)
            return
        
        self.cubes = []
        size = min(width, height) // 10
        
        for i, (x, y) in enumerate(self.positions):
            px = (x / 100) * width
            py = (y / 100) * height
            
            # Sombra
            shadow = self.board_canvas.create_rectangle(
                px - size//2 + 3, py - size//2 + 3,
                px + size//2 + 3, py + size//2 + 3,
                fill='#0a0a0a', outline=''
            )
            
            # Cubo
            cube_id = self.round_rect(px - size//2, py - size//2,
                                     px + size//2, py + size//2,
                                     12, fill='#1e293b', outline='#475569', width=2)
            
            # Número
            text_id = self.board_canvas.create_text(px, py, text=str(i+1),
                                                   font=('Segoe UI', 18, 'bold'),
                                                   fill='#64748b')
            
            # Zona clickeable
            hit = self.board_canvas.create_rectangle(
                px - size//2, py - size//2, px + size//2, py + size//2,
                fill='', outline='', tags=f'hit_{i}'
            )
            self.board_canvas.tag_bind(f'hit_{i}', '<Button-1>',
                                      lambda e, idx=i: self.on_cube_click(idx))
            
            self.cubes.append({
                'idx': i, 'x': px, 'y': py, 'size': size,
                'cube_id': cube_id, 'text_id': text_id, 'shadow_id': shadow
            })
    
    def round_rect(self, x1, y1, x2, y2, radius, **kwargs):
        points = [x1+radius, y1, x2-radius, y1, x2, y1, x2, y1+radius,
                 x2, y2-radius, x2, y2, x2-radius, y2, x1+radius, y2,
                 x1, y2, x1, y2-radius, x1, y1+radius, x1, y1]
        return self.board_canvas.create_polygon(points, smooth=True, **kwargs)
    
    def set_cube(self, idx, state):
        """Cambiar estado del cubo"""
        if not self.cubes or idx >= len(self.cubes):
            return
        
        c = self.cubes[idx]
        
        colors = {
            'default': ('#1e293b', '#64748b', '#475569'),
            'active': (self.colors['accent'], '#0f172a', '#00d4ff'),
            'selected': (self.colors['success'], 'white', '#10b981'),
            'error': (self.colors['error'], 'white', '#ef4444'),
            'disabled': ('#0f172a', '#1e293b', '#1e293b')
        }
        
        bg, fg, border = colors.get(state, colors['default'])
        
        try:
            self.board_canvas.itemconfig(c['cube_id'], fill=bg, outline=border)
            self.board_canvas.itemconfig(c['text_id'], fill=fg)
        except:
            pass
    
    def start_level(self):
        """Iniciar nivel"""
        if self.level_in_progress:
            return
        
        self.level_in_progress = True
        self.first_level_started = True
        
        # Cambiar color visual del botón
        self.btn_start.itemconfig(1, fill='#4a5568')
        
        self.clear_selection()
        self.sequence = random.sample(range(9), self.current_level)
        self.level_start_time = time.time()
        
        self.update_display()
        self.show_message(f"Nivel {self.current_level}: Observa la secuencia...", 'info')
        
        for i in range(9):
            self.set_cube(i, 'disabled')
        
        # Delay inicial más largo para asegurar que el sistema de audio esté listo
        initial_delay = 1500  # Aumentado a 1.5 segundos para asegurar inicialización
        self.safe_after(initial_delay, self.show_sequence)
    
    def show_sequence(self):
        """Mostrar secuencia"""
        self.is_showing = True
        self.sequence_start_time = time.time()
        
        self.log_event('sequence_start', None, {
            'level': self.current_level,
            'sequence_length': len(self.sequence)
        })
        
        # Pequeño delay adicional antes del primer beep para asegurar que todo esté listo
        self.safe_after(200, lambda: self.animate_cube(0))
    
    def animate_cube(self, idx):
        """Animar cubo"""
        if idx < len(self.sequence):
            c = self.sequence[idx]
            self.set_cube(c, 'active')
            
            # Sonido con manejo robusto de errores
            self.play_beep_safe(600 + idx*50)
            
            self.safe_after(700, lambda c=c, i=idx: self.hide_cube(c, i))
        else:
            self.is_showing = False
            self.can_click = True
            self.last_click_time = time.time()
            
            for i in range(9):
                self.set_cube(i, 'default')
            
            msg = "Repite el MISMO orden" if self.test_mode == 'normal' else "Repite al REVÉS"
            self.show_message(msg, 'success')
    
    def play_beep_safe(self, frequency, duration=180):
        """Reproducir beep de forma segura con manejo de errores"""
        try:
            import winsound
            # Usar MB_OK en lugar de Beep si hay problemas, o intentar Beep con manejo de excepción
            winsound.Beep(frequency, duration)
        except Exception as e:
            # Si falla, intentar con MessageBeep como fallback
            try:
                import winsound
                winsound.MessageBeep(winsound.MB_OK)
            except:
                pass  # Silencio si todo falla
    
    def hide_cube(self, cube_id, idx):
        """Ocultar cubo"""
        self.set_cube(cube_id, 'disabled')
        self.safe_after(250, lambda: self.animate_cube(idx + 1))
    
    def repeat_sequence(self):
        """Repetir secuencia"""
        if not self.level_in_progress or not self.sequence or self.is_showing:
            return
        
        self.can_click = False
        self.clear_selection()
        self.show_message("Repitiendo secuencia...", 'warning')
        
        for i in range(9):
            self.set_cube(i, 'disabled')
        
        self.safe_after(400, self.show_sequence)
    
    def on_cube_click(self, cube_id):
        """Click en cubo"""
        if not self.can_click or self.is_showing or not self.level_in_progress:
            return
        
        if cube_id in self.user_sequence:
            return
        
        now = time.time()
        reaction_time = (now - self.sequence_start_time) * 1000 if not self.user_sequence else (now - self.last_click_time) * 1000
        
        # Calcular esperado
        if self.test_mode == 'normal':
            expected = self.sequence[len(self.user_sequence)]
        else:
            expected = self.sequence[len(self.sequence) - 1 - len(self.user_sequence)]
        
        correct = (cube_id == expected)
        
        self.log_event('cube_click', cube_id, {
            'sequence_position': len(self.user_sequence),
            'expected_cube': expected,
            'is_correct': correct,
            'reaction_time_ms': reaction_time,
            'level': self.current_level
        })
        
        if correct:
            self.set_cube(cube_id, 'selected')
            self.user_sequence.append(cube_id)
            self.last_click_time = now
            self.play_beep_safe(800, 150)  # Beep más corto para respuesta correcta
            
            if len(self.user_sequence) == len(self.sequence):
                self.handle_success()
        else:
            self.set_cube(cube_id, 'error')
            self.play_beep_safe(300, 300)  # Beep más largo para error
            self.log_event('error', cube_id, {
                'expected_cube': expected,
                'error_type': 'commission'
            })
            self.safe_after(400, self.handle_error)
    
    def handle_success(self):
        """Nivel exitoso"""
        self.can_click = False
        level_time = (time.time() - self.level_start_time) * 1000
        
        if self.current_level > self.corsi_span:
            self.corsi_span = self.current_level
        
        # Calcular métricas del nivel
        level_clicks = [m for m in self.movements 
                       if m.get('level') == self.current_level and m.get('event') == 'cube_click']
        reaction_times = [m.get('reaction_time_ms', 0) for m in level_clicks if m.get('reaction_time_ms')]
        
        avg_reaction = statistics.mean(reaction_times) if reaction_times else 0
        first_reaction = reaction_times[0] if reaction_times else 0
        
        self.summaries.append({
            'level': self.current_level,
            'success': True,
            'attempts': 3 - self.attempts_left,
            'sequence_length': self.current_level,
            'total_time_ms': level_time,
            'avg_reaction_time_ms': avg_reaction,
            'first_reaction_time_ms': first_reaction,
            'error_count': 0
        })
        
        self.show_message("¡Correcto! ✅", 'success')
        self.level_in_progress = False
        self.metadata['total_levels_success'] = len([s for s in self.summaries if s['success']])
        self.metadata['total_levels_attempted'] = len(self.summaries)
        
        if self.current_level < 9:
            self.current_level += 1
            self.attempts_left = 2
            self.update_display()
            self.show_message(f"Nivel completado. Presiona INICIAR para el nivel {self.current_level}", 'success')
            # Resetear botón
            self.btn_start.itemconfig(1, fill=self.colors['accent'])
        else:
            self.show_message("¡Prueba completada!", 'success')
            self.safe_after(1500, self.finish)
    
    def handle_error(self):
        """Error en nivel"""
        self.attempts_left -= 1
        self.update_display()
        
        if self.attempts_left > 0:
            self.show_message(f"Error ❌ Quedan {self.attempts_left} intentos", 'error')
            self.clear_selection()
            self.safe_after(1200, self.show_sequence)
        else:
            level_time = (time.time() - self.level_start_time) * 1000
            
            self.summaries.append({
                'level': self.current_level,
                'success': False,
                'attempts': 2,
                'sequence_length': self.current_level,
                'total_time_ms': level_time,
                'error_count': 2
            })
            
            self.show_message(f"Nivel {self.current_level} fallido", 'error')
            self.level_in_progress = False
            self.metadata['total_levels_attempted'] = len(self.summaries)
            self.safe_after(1500, self.finish)
    
    def clear_selection(self):
        """Limpiar selección"""
        self.user_sequence = []
        for i in range(9):
            self.set_cube(i, 'default')
    
    def update_display(self):
        """Actualizar display"""
        self.lbl_level.config(text=str(self.current_level))
        self.lbl_attempts.config(text=str(self.attempts_left))
        self.lbl_span.config(text=str(self.corsi_span) if self.corsi_span > 0 else "-")
    
    def show_message(self, text, type):
        """Mostrar mensaje"""
        colors = {
            'info': '#60a5fa',
            'success': '#10b981',
            'warning': '#f59e0b',
            'error': '#ef4444'
        }
        self.lbl_message.config(text=text, fg=colors.get(type, 'white'))
    
    def log_event(self, event, cube, data):
        """Registrar evento"""
        self.movements.append({
            'timestamp': datetime.now().isoformat(),
            'event': event,
            'cube': cube,
            'level': self.current_level,
            **data
        })
    
    def beep(self, freq):
        """Sonido - mantenido para compatibilidad pero usa el nuevo método"""
        self.play_beep_safe(freq)
    
    def confirm_stop(self):
        """Confirmar terminación"""
        if not self.level_in_progress and not self.summaries:
            if messagebox.askyesno("Salir", "¿Deseas salir sin guardar?"):
                self.show_welcome()
            return
        
        if messagebox.askyesno("Terminar", "¿Terminar prueba?\nSe guardarán los datos."):
            self.finish()
    
    def finish(self):
        """Finalizar y guardar"""
        self.metadata['final_span'] = self.corsi_span
        self.save_data()
        self.show_post_test_options()
    
    def save_data(self):
        """Guardar datos"""
        try:
            # Metadata CSV
            with open(os.path.join(DATA_DIR, f'{self.session_id}_metadata.csv'), 'w', newline='', encoding='utf-8') as f:
                w = csv.DictWriter(f, fieldnames=self.metadata.keys())
                w.writeheader()
                w.writerow(self.metadata)
            
            # Events CSV
            if self.movements:
                all_fields = set()
                for m in self.movements:
                    all_fields.update(m.keys())
                
                with open(os.path.join(DATA_DIR, f'{self.session_id}_events.csv'), 'w', newline='', encoding='utf-8') as f:
                    w = csv.DictWriter(f, fieldnames=sorted(all_fields))
                    w.writeheader()
                    for m in self.movements:
                        row = {k: m.get(k, '') for k in all_fields}
                        w.writerow(row)
            
            # Excel completo
            self.create_excel()
            
        except Exception as e:
            messagebox.showerror("Error guardando", str(e))
            print(f"Error detallado: {e}")
    
    def create_excel(self):
        """Crear Excel estructurado similar al D2"""
        self.excel_path = os.path.join(DATA_DIR, f'{self.session_id}_completo.xlsx')
        wb = Workbook()
        
        # Estilos
        header_fill = PatternFill(start_color="0ea5e9", end_color="0ea5e9", fill_type="solid")
        header_font = Font(color="ffffff", bold=True, size=11)
        title_font = Font(size=16, bold=True, color="0ea5e9")
        section_font = Font(size=12, bold=True, color="64748b")
        border = Border(left=Side(style='thin'), right=Side(style='thin'),
                       top=Side(style='thin'), bottom=Side(style='thin'))
        
        # HOJA 1: INFO
        ws_info = wb.active
        ws_info.title = "Info"
        
        ws_info['A1'] = "CORSI - INFORMACIÓN DEL PARTICIPANTE"
        ws_info['A1'].font = title_font
        
        info_data = [
            ("ID Sesión", self.session_id),
            ("ID Participante", self.metadata['user_id']),
            ("Nombre", self.metadata.get('user_name', 'No especificado')),
            ("Edad", self.metadata['age']),
            ("Género", self.metadata['gender']),
            ("Condición", self.metadata['condition']),
            ("Modo de Prueba", "Directo" if self.test_mode == 'normal' else "Inverso"),
            ("Fecha", self.metadata['timestamp'][:10]),
            ("Hora", self.metadata['timestamp'][11:19]),
            ("", ""),
            ("RESULTADOS", ""),
            ("Span Final", self.corsi_span),
            ("Nivel Alcanzado", self.current_level),
            ("Niveles Completados", len([s for s in self.summaries if s['success']])),
            ("Niveles Intentados", len(self.summaries)),
        ]
        
        for i, (k, v) in enumerate(info_data, 4):
            if k == "RESULTADOS":
                ws_info.cell(row=i, column=1, value=k).font = section_font
            else:
                ws_info.cell(row=i, column=1, value=k).font = Font(bold=True)
                ws_info.cell(row=i, column=2, value=v)
        
        ws_info.column_dimensions['A'].width = 25
        ws_info.column_dimensions['B'].width = 40
        
        # HOJA 2: Métricas
        ws_metrics = wb.create_sheet("Metricas")
        
        headers = ['Nivel', 'Éxito', 'Intentos', 'Longitud_Sec', 'Tiempo_Total_ms', 
                  'Tiempo_Promedio_ms', 'Primera_Reaccion_ms', 'Errores']
        
        for col, h in enumerate(headers, 1):
            cell = ws_metrics.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
        
        for row_idx, s in enumerate(self.summaries, 2):
            ws_metrics.cell(row=row_idx, column=1, value=s['level'])
            
            success = s['success']
            cell = ws_metrics.cell(row=row_idx, column=2, value='Sí' if success else 'No')
            cell.fill = PatternFill(start_color="C6EFCE" if success else "FFC7CE", fill_type="solid")
            
            ws_metrics.cell(row=row_idx, column=3, value=s['attempts'])
            ws_metrics.cell(row=row_idx, column=4, value=s['sequence_length'])
            ws_metrics.cell(row=row_idx, column=5, value=round(s.get('total_time_ms', 0), 2))
            ws_metrics.cell(row=row_idx, column=6, value=round(s.get('avg_reaction_time_ms', 0), 2))
            ws_metrics.cell(row=row_idx, column=7, value=round(s.get('first_reaction_time_ms', 0), 2))
            ws_metrics.cell(row=row_idx, column=8, value=s.get('error_count', 0))
            
            for c in range(1, 9):
                ws_metrics.cell(row=row_idx, column=c).border = border
        
        for col in range(1, 9):
            ws_metrics.column_dimensions[chr(64+col)].width = 18
        
        # HOJA 3: Eventos
        ws_events = wb.create_sheet("Eventos")
        
        event_headers = ['Timestamp', 'Evento', 'Nivel', 'Cubo', 'Posicion_Sec', 
                        'Cubo_Esperado', 'Correcto', 'Tiempo_Reaccion_ms']
        
        for col, h in enumerate(event_headers, 1):
            cell = ws_events.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
        
        for row_idx, m in enumerate(self.movements, 2):
            ws_events.cell(row=row_idx, column=1, value=str(m.get('timestamp', ''))[:19])
            ws_events.cell(row=row_idx, column=2, value=m.get('event', ''))
            ws_events.cell(row=row_idx, column=3, value=m.get('level', ''))
            ws_events.cell(row=row_idx, column=4, value=m.get('cube', ''))
            ws_events.cell(row=row_idx, column=5, value=m.get('sequence_position', ''))
            ws_events.cell(row=row_idx, column=6, value=m.get('expected_cube', ''))
            
            correct = m.get('is_correct')
            cell = ws_events.cell(row=row_idx, column=7, value='Sí' if correct else ('No' if correct == False else ''))
            if correct == True:
                cell.fill = PatternFill(start_color="C6EFCE", fill_type="solid")
            elif correct == False:
                cell.fill = PatternFill(start_color="FFC7CE", fill_type="solid")
            
            rt = m.get('reaction_time_ms')
            ws_events.cell(row=row_idx, column=8, value=round(rt, 2) if rt else '')
            
            for c in range(1, 9):
                ws_events.cell(row=row_idx, column=c).border = border
        
        for col in range(1, 9):
            ws_events.column_dimensions[chr(64+col)].width = 16
        
        # HOJA 4: ML_Features
        ws_ml = wb.create_sheet("ML_Features")
        
        ml_features = self.calculate_ml_features()
        
        ml_headers = list(ml_features.keys())
        
        for col, h in enumerate(ml_headers, 1):
            cell = ws_ml.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
        
        for col, (k, v) in enumerate(ml_features.items(), 1):
            ws_ml.cell(row=2, column=col, value=v).border = border
        
        for col in range(1, len(ml_headers) + 1):
            ws_ml.column_dimensions[chr(64+col) if col <= 26 else 'A'+chr(64+col-26)].width = 20
        
        wb.save(self.excel_path)
    
    def calculate_ml_features(self):
        """Calcular features para ML"""
        features = {
            'session_id': self.session_id,
            'user_id': self.metadata['user_id'],
            'user_name': self.metadata.get('user_name', ''),
            'age': self.metadata['age'],
            'gender': self.metadata['gender'],
            'condition': self.metadata['condition'],
            'test_mode': 0 if self.test_mode == 'normal' else 1,
            
            'final_span': self.corsi_span,
            'max_level_reached': self.current_level,
            'total_levels_completed': len([s for s in self.summaries if s['success']]),
            'total_levels_attempted': len(self.summaries),
            'success_rate': len([s for s in self.summaries if s['success']]) / len(self.summaries) if self.summaries else 0,
            
            'total_test_time_ms': sum([s.get('total_time_ms', 0) for s in self.summaries]),
        }
        
        all_reaction_times = []
        for m in self.movements:
            if m.get('event') == 'cube_click' and m.get('reaction_time_ms'):
                all_reaction_times.append(m.get('reaction_time_ms'))
        
        if all_reaction_times:
            features['mean_reaction_time_ms'] = statistics.mean(all_reaction_times)
            features['std_reaction_time_ms'] = statistics.stdev(all_reaction_times) if len(all_reaction_times) > 1 else 0
            features['min_reaction_time_ms'] = min(all_reaction_times)
            features['max_reaction_time_ms'] = max(all_reaction_times)
            features['first_reaction_time_ms'] = all_reaction_times[0] if all_reaction_times else 0
        else:
            features['mean_reaction_time_ms'] = 0
            features['std_reaction_time_ms'] = 0
            features['min_reaction_time_ms'] = 0
            features['max_reaction_time_ms'] = 0
            features['first_reaction_time_ms'] = 0
        
        for level in range(2, 7):
            level_data = [s for s in self.summaries if s['level'] == level]
            if level_data:
                s = level_data[0]
                features[f'level_{level}_success'] = 1 if s['success'] else 0
                features[f'level_{level}_time_ms'] = s.get('total_time_ms', 0)
                features[f'level_{level}_avg_rt_ms'] = s.get('avg_reaction_time_ms', 0)
            else:
                features[f'level_{level}_success'] = -1
                features[f'level_{level}_time_ms'] = 0
                features[f'level_{level}_avg_rt_ms'] = 0
        
        error_events = [m for m in self.movements if m.get('event') == 'error']
        features['total_errors'] = len(error_events)
        features['error_rate'] = len(error_events) / len([m for m in self.movements if m.get('event') == 'cube_click']) if self.movements else 0
        
        if self.test_mode == 'inverse':
            features['inverse_mode'] = 1
            features['cognitive_load_proxy'] = features['mean_reaction_time_ms'] * features['total_errors']
        else:
            features['inverse_mode'] = 0
            features['cognitive_load_proxy'] = 0
        
        return features
    
    def show_post_test_options(self):
        """Pantalla de opciones post-prueba - CORREGIDA"""
        self.clear()
        
        # Header
        header = tk.Frame(self.root, bg='#252545', height=100)
        header.pack(fill=tk.X)
        header.pack_propagate(False)
        
        header_content = tk.Frame(header, bg='#252545')
        header_content.place(relx=0.5, rely=0.5, anchor='center')
        
        tk.Label(header_content, text="✅ Prueba Finalizada", 
                font=('Segoe UI', 24, 'bold'), fg='#10b981', bg='#252545').pack()
        
        # Contenido
        main = tk.Frame(self.root, bg=self.colors['bg'])
        main.pack(fill=tk.BOTH, expand=True, padx=80, pady=30)
        
        card = tk.Frame(main, bg='#2d2a4a', highlightbackground='#4a5568',
                       highlightthickness=1, padx=40, pady=30)
        card.pack(fill=tk.BOTH, expand=True)
        
        # Resumen rápido
        summary_frame = tk.Frame(card, bg='#1a1c2e', padx=20, pady=15)
        summary_frame.pack(fill=tk.X, pady=(0, 20))
        
        mode_text = "Directo" if self.test_mode == 'normal' else "Inverso"
        tk.Label(summary_frame, text=f"Modo: {mode_text}  |  Span: {self.corsi_span}  |  Nivel: {self.current_level}", 
                font=('Segoe UI', 12), fg='white', bg='#1a1c2e').pack()
        
        # Pregunta principal
        tk.Label(card, text="¿Qué deseas hacer ahora?", 
                font=('Segoe UI', 18, 'bold'), fg=self.colors['accent'], 
                bg='#2d2a4a').pack(pady=20)
        
        # Opciones en grid 3x2 (incluyendo salir)
        options_frame = tk.Frame(card, bg='#2d2a4a')
        options_frame.pack(pady=10)
        
        # Opción 1: Misma prueba, otro participante
        btn_same = self.create_rounded_button(options_frame, 
                                             f"🔄 Repetir {mode_text}\n(mismo modo)", 
                                             self.restart_same_mode,
                                             self.colors['accent'], '#0f172a',
                                             width=220, height=70, font_size=11)
        btn_same.grid(row=0, column=0, padx=10, pady=10)
        
        # Opción 2: Cambiar modo
        other_mode = "Inverso" if self.test_mode == 'normal' else "Directo"
        other_color = self.colors['warning'] if self.test_mode == 'normal' else self.colors['accent']
        btn_switch = self.create_rounded_button(options_frame, 
                                               f"🔄 Cambiar a {other_mode}\n(otro modo)", 
                                               self.switch_mode,
                                               other_color, '#0f172a',
                                               width=220, height=70, font_size=11)
        btn_switch.grid(row=0, column=1, padx=10, pady=10)
        
        # Opción 3: Ver historial
        btn_history = self.create_rounded_button(options_frame, 
                                                "📂 Ver Historial\nde Pruebas", 
                                                self.show_history,
                                                '#4a5568', 'white',
                                                width=220, height=70, font_size=11)
        btn_history.grid(row=1, column=0, padx=10, pady=10)
        
        # Opción 4: Menú principal
        btn_menu = self.create_rounded_button(options_frame, 
                                             "🏠 Menú Principal\n(inicio)", 
                                             self.show_welcome,
                                             '#4a5568', 'white',
                                             width=220, height=70, font_size=11)
        btn_menu.grid(row=1, column=1, padx=10, pady=10)
        
        # Opción 5: Salir (ahora visible en el grid)
        btn_exit = self.create_rounded_button(options_frame, 
                                             "❌ Salir del\nPrograma", 
                                             self.root.quit,
                                             '#ef4444', 'white',
                                             width=220, height=70, font_size=11)
        btn_exit.grid(row=2, column=0, columnspan=2, padx=10, pady=10)
        
        # Archivo generado con botón de abrir (debajo del grid)
        file_frame = tk.Frame(card, bg='#2d2a4a')
        file_frame.pack(fill=tk.X, pady=(20, 10))
        
        tk.Label(file_frame, text="📁 Archivo Excel generado:", 
                font=('Segoe UI', 11), fg='#94a3b8', bg='#2d2a4a').pack()
        
        filename = os.path.basename(self.excel_path) if self.excel_path else "Error"
        tk.Label(file_frame, text=filename,
                font=('Segoe UI', 10), fg='#60a5fa', bg='#2d2a4a',
                wraplength=600).pack(pady=(5, 10))
        
        # Botón abrir Excel
        if self.excel_path and os.path.exists(self.excel_path):
            btn_open = self.create_rounded_button(file_frame, "📊 Abrir Excel", 
                                                 self.open_excel,
                                                 '#10b981', 'white',
                                                 width=180, height=40, font_size=11)
            btn_open.pack(pady=5)
    
    def open_excel(self):
        """Abrir el archivo Excel generado"""
        if self.excel_path and os.path.exists(self.excel_path):
            try:
                os.startfile(self.excel_path)
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo abrir el archivo:\n{str(e)}")
        else:
            messagebox.showerror("Error", "Archivo no encontrado")
    
    def restart_same_mode(self):
        """Reiniciar con el mismo modo"""
        self.reset_session()
        self.show_config()
    
    def switch_mode(self):
        """Cambiar al modo opuesto"""
        self.reset_session()
        self.test_mode = 'inverse' if self.test_mode == 'normal' else 'normal'
        self.show_config()
    
    def reset_session(self):
        """Resetear variables de sesión"""
        self.current_level = 2
        self.attempts_left = 2
        self.corsi_span = 0
        self.sequence = []
        self.user_sequence = []
        self.movements = []
        self.summaries = []
        self.metadata = {}
        self.session_id = ''
        self.level_in_progress = False
        self.first_level_started = False
        self.excel_path = None
        if hasattr(self, '_first_sound_played'):
            delattr(self, '_first_sound_played')
    
    def show_history(self):
        """Mostrar historial de pruebas - similar a D2"""
        self.clear()
        
        # Header
        header = tk.Frame(self.root, bg='#252545', height=100)
        header.pack(fill=tk.X)
        header.pack_propagate(False)
        
        header_content = tk.Frame(header, bg='#252545')
        header_content.place(relx=0.5, rely=0.5, anchor='center')
        
        tk.Label(header_content, text="📂 Historial de Pruebas", 
                font=('Segoe UI', 24, 'bold'), fg='#60a5fa', bg='#252545').pack()
        
        # Contenido
        main = tk.Frame(self.root, bg=self.colors['bg'])
        main.pack(fill=tk.BOTH, expand=True, padx=40, pady=20)
        
        card = tk.Frame(main, bg='#2d2a4a', highlightbackground='#4a5568',
                       highlightthickness=1, padx=20, pady=20)
        card.pack(fill=tk.BOTH, expand=True)
        
        # Botón volver
        btn_back = self.create_rounded_button(card, "← Volver", 
                                             self.show_welcome,
                                             '#4a5568', 'white', 
                                             width=100, height=35, font_size=10)
        btn_back.pack(anchor='nw', pady=(0, 10))
        
        # Treeview para historial
        columns = ('Fecha', 'Hora', 'ID', 'Nombre', 'Modo', 'Span', 'Nivel', 'Archivo')
        tree = ttk.Treeview(card, columns=columns, show='headings', height=15)
        
        for col in columns:
            tree.heading(col, text=col)
            tree.column(col, width=100 if col != 'Nombre' else 150)
        
        # Scrollbar
        scrollbar = ttk.Scrollbar(card, orient=tk.VERTICAL, command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        
        # Empaquetar
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Cargar archivos
        try:
            files = [f for f in os.listdir(DATA_DIR) if f.endswith('_completo.xlsx')]
            files.sort(reverse=True)  # Más recientes primero
            
            for file in files:
                parts = file.replace('_completo.xlsx', '').split('_')
                if len(parts) >= 3:
                    date_str = parts[1]
                    time_str = parts[2]
                    fecha = f"{date_str[:4]}-{date_str[4:6]}-{date_str[6:]}"
                    hora = f"{time_str[:2]}:{time_str[2:4]}:{time_str[4:]}"
                    
                    # Intentar leer metadata
                    metadata_file = file.replace('_completo.xlsx', '_metadata.csv')
                    metadata_path = os.path.join(DATA_DIR, metadata_file)
                    
                    nombre = "Desconocido"
                    modo = "?"
                    span = "?"
                    nivel = "?"
                    user_id = "?"
                    
                    if os.path.exists(metadata_path):
                        try:
                            with open(metadata_path, 'r', encoding='utf-8') as f:
                                reader = csv.DictReader(f)
                                meta = next(reader, {})
                                nombre = meta.get('user_name', 'Desconocido')[:15]
                                user_id = meta.get('user_id', '?')
                                modo_raw = meta.get('test_mode', '')
                                modo = "Directo" if modo_raw == 'normal' else "Inverso" if modo_raw == 'inverse' else "?"
                                span = meta.get('final_span', '?')
                                nivel = meta.get('total_levels_attempted', '?')
                        except:
                            pass
                    
                    tree.insert('', tk.END, values=(fecha, hora, user_id, nombre, modo, span, nivel, file))
        except Exception as e:
            messagebox.showerror("Error", f"Error cargando historial: {str(e)}")
        
        # Botón abrir seleccionado
        def open_selected():
            selected = tree.selection()
            if not selected:
                messagebox.showwarning("Seleccionar", "Selecciona un archivo primero")
                return
            
            item = tree.item(selected[0])
            filename = item['values'][-1]
            filepath = os.path.join(DATA_DIR, filename)
            
            if os.path.exists(filepath):
                try:
                    os.startfile(filepath)
                except Exception as e:
                    messagebox.showerror("Error", f"No se pudo abrir: {str(e)}")
            else:
                messagebox.showerror("Error", "Archivo no encontrado")
        
        btn_open = self.create_rounded_button(card, "📊 Abrir Excel Seleccionado", 
                                             open_selected,
                                             '#10b981', 'white',
                                             width=220, height=40, font_size=11)
        btn_open.pack(pady=(10, 0))
    
    def restart(self):
        """Reiniciar completo (botón legacy)"""
        self.reset_session()
        self.test_mode = ''
        self.show_welcome()


if __name__ == '__main__':
    root = tk.Tk()
    app = CorsiApp(root)
    root.mainloop()